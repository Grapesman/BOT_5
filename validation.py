from dataclasses import dataclass
from datetime import datetime
from typing import Any, Callable, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from logger import logger


def has_letters(value: Any) -> bool:
    return isinstance(value, str) and any(char.isalpha() for char in value)


def is_date(value: Any) -> bool:
    return isinstance(value, datetime)


def is_string(value: Any) -> bool:
    return isinstance(value, str)


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


@dataclass(frozen=True)
class ColumnRule:
    """Правило для заполненной ячейки: пустая ячейка считается допустимой"""
    column: str
    name: str
    expected: str
    is_valid: Callable[[Any], bool]

    def check(self, value: Any) -> Optional[str]:
        if value is None or self.is_valid(value):
            return None
        return f"столбец {self.column} ({self.name}) — ожидается {self.expected}, получено {value!r}"


class SheetValidator:
    """Удаляет из листа строки, не прошедшие проверку, и пустой хвост после последней строки с данными"""

    def __init__(self, sheet_name: str, first_row: int, rules: List[ColumnRule]):
        self.sheet_name = sheet_name
        self.first_row = first_row
        self.rules = rules

    def validate(self, book: Workbook) -> None:
        sheet = book[self.sheet_name]
        self._trim_empty_tail(sheet)

        invalid_rows = self._find_invalid_rows(sheet)
        for row, reasons in invalid_rows.items():
            logger.warning(f"Лист '{self.sheet_name}', строка {row} удалена: {'; '.join(reasons)}")
        for row in sorted(invalid_rows, reverse=True):
            sheet.delete_rows(row)

        remaining = max(sheet.max_row - self.first_row + 1, 0)
        logger.info(f"Валидация листа '{self.sheet_name}': удалено строк — {len(invalid_rows)}, осталось — {remaining}")

    def _find_invalid_rows(self, sheet: Worksheet) -> Dict[int, List[str]]:
        invalid_rows = {}
        for row in range(self.first_row, sheet.max_row + 1):
            reasons = [
                reason for rule in self.rules
                if (reason := rule.check(sheet[f"{rule.column}{row}"].value))
            ]
            if reasons:
                invalid_rows[row] = reasons
        return invalid_rows

    def last_data_row(self, sheet: Worksheet) -> int:
        for row in range(sheet.max_row, self.first_row - 1, -1):
            if not all(is_blank(sheet[f"{rule.column}{row}"].value) for rule in self.rules):
                return row
        return self.first_row - 1

    def _trim_empty_tail(self, sheet: Worksheet) -> None:
        last_row = self.last_data_row(sheet)
        if sheet.max_row > last_row:
            sheet.delete_rows(last_row + 1, sheet.max_row - last_row)


catalog_validator = SheetValidator(
    sheet_name="Каталог статей",
    first_row=2,
    rules=[
        ColumnRule("B", "Название", "текст с буквами", has_letters),
        ColumnRule("C", "Дата создания", "дата", is_date),
        ColumnRule("D", "Дата публикации", "дата", is_date),
        ColumnRule("H", "Авторы", "текст с буквами", has_letters),
        ColumnRule("J", "Ключевые слова", "строка", is_string),
    ],
)
